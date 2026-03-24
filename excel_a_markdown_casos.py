from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import OrderedDict
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


DEFAULT_JSON_OUTPUT = "catalogo_casos_prueba.json"
DEFAULT_MARKDOWN_OUTPUT = "catalogo_casos_prueba.md"
DEFAULT_PREFIX_REPORT = "reporte_asociacion_por_prefijo.md"
SCENARIOS_SHEET = "ESCENARIOS"
TEST_CASES_SHEET = "Casos de prueba"


def slugify_header(value: object) -> str:
    text = normalize_scalar(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.casefold()
    text = re.sub(r"[^a-z0-9]+", "", text)
    return text


def normalize_scalar(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def clean_inline_text(value: object) -> str:
    text = normalize_scalar(value)
    return re.sub(r"\s+", " ", text).strip()


def split_nonempty_lines(value: object) -> list[str]:
    text = normalize_scalar(value)
    if not text:
        return []
    return [line.strip() for line in text.split("\n") if line.strip()]


def normalize_multiline_text(value: object) -> str:
    lines = split_nonempty_lines(value)
    return "\n".join(lines)


def normalize_numbered_items(value: object) -> list[str]:
    lines = split_nonempty_lines(value)
    items: list[str] = []
    pattern = re.compile(r"^\s*(\d+[\.\-\)]|[-*])\s*(.+?)\s*$")

    for line in lines:
        match = pattern.match(line)
        items.append(match.group(2).strip() if match else line)

    if not items and clean_inline_text(value):
        items = [clean_inline_text(value)]
    return items


def normalize_bullet_items(value: object) -> list[str]:
    lines = split_nonempty_lines(value)
    items: list[str] = []
    pattern = re.compile(r"^\s*[-*]\s*(.+?)\s*$")

    for line in lines:
        match = pattern.match(line)
        items.append(match.group(1).strip() if match else line)

    if not items and clean_inline_text(value):
        items = [clean_inline_text(value)]
    return items


def find_sheet_name(sheet_names: Iterable[str], expected_name: str) -> str:
    expected = expected_name.casefold()
    for name in sheet_names:
        if name.casefold() == expected:
            return name
    raise ValueError(f"No se encontro la hoja '{expected_name}' en el archivo Excel.")


def find_scenarios_header_row(worksheet) -> int:
    for row_index, row in enumerate(worksheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        slugs = [slugify_header(cell) for cell in row]
        if "id" in slugs and "escenariosdetectados" in slugs:
            return row_index
    raise ValueError("No se encontro la fila de encabezados valida en la hoja ESCENARIOS.")


def build_header_index(row: tuple[object, ...]) -> dict[str, int]:
    header_index: dict[str, int] = {}
    for column_index, cell in enumerate(row):
        slug = slugify_header(cell)
        if slug:
            header_index[slug] = column_index
    return header_index


def load_scenarios(workbook_path: Path) -> OrderedDict[str, dict]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    try:
        sheet_name = find_sheet_name(workbook.sheetnames, SCENARIOS_SHEET)
        worksheet = workbook[sheet_name]

        header_row_index = find_scenarios_header_row(worksheet)
        header_row = next(
            worksheet.iter_rows(min_row=header_row_index, max_row=header_row_index, values_only=True)
        )
        header_index = build_header_index(header_row)

        n_col = header_index.get("n")
        id_col = header_index.get("id")
        description_col = header_index.get("escenariosdetectados")
        if n_col is None:
            raise ValueError("No se encontro la columna 'N' en la hoja ESCENARIOS.")
        if id_col is None:
            raise ValueError("No se encontro la columna 'ID' en la hoja ESCENARIOS.")
        if description_col is None:
            raise ValueError(
                "No se encontro la columna 'Escenarios detectados' en la hoja ESCENARIOS."
            )

        scenarios: OrderedDict[str, dict] = OrderedDict()
        visible_sequence_numbers: list[int] = []
        for row_index, row in enumerate(
            worksheet.iter_rows(min_row=header_row_index + 1, values_only=True),
            start=header_row_index + 1,
        ):
            if worksheet.row_dimensions[row_index].hidden:
                continue
            scenario_id = clean_inline_text(row[id_col] if id_col < len(row) else "")
            description = clean_inline_text(row[description_col] if description_col < len(row) else "")
            sequence_value = row[n_col] if n_col < len(row) else None
            if not scenario_id:
                continue

            if sequence_value is not None:
                try:
                    visible_sequence_numbers.append(int(sequence_value))
                except (TypeError, ValueError):
                    raise ValueError(
                        f"Valor no numerico en la columna 'N' de ESCENARIOS para la fila {row_index}: {sequence_value!r}"
                    )

            scenarios[scenario_id] = {
                "id": scenario_id,
                "descripcion": description,
                "descripcion_normalizada": normalize_lookup_text(description),
                "casos": [],
            }

        if not scenarios:
            raise ValueError("No se encontraron escenarios validos en la hoja ESCENARIOS.")

        if visible_sequence_numbers:
            expected_sequence = list(range(min(visible_sequence_numbers), max(visible_sequence_numbers) + 1))
            missing_values = [value for value in expected_sequence if value not in visible_sequence_numbers]
            if missing_values:
                print(
                    "Advertencia: la secuencia visible de la columna 'N' en ESCENARIOS tiene huecos: "
                    + ", ".join(str(value) for value in missing_values),
                    file=sys.stderr,
                )
        return scenarios
    finally:
        workbook.close()


def load_test_cases(workbook_path: Path) -> list[dict]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet_name = find_sheet_name(workbook.sheetnames, TEST_CASES_SHEET)
        worksheet = workbook[sheet_name]

        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        header_index = build_header_index(header_row)

        required_headers = {
            "historiadeusuariorequerimientoescenario": "Historia de usuario/ Requerimiento/ Escenario",
            "titulo": "Titulo/Título",
            "descripcion": "Descripcion/Descripción",
            "precondiciones": "Precondiciones",
            "datos": "Datos",
            "pasos": "Pasos",
            "resultadoesperado": "Resultado Esperado",
        }

        missing_headers = [
            label for slug, label in required_headers.items() if slug not in header_index
        ]
        if missing_headers:
            raise ValueError(
                "Faltan columnas minimas en la hoja 'Casos de prueba': "
                + ", ".join(missing_headers)
            )

        test_cases: list[dict] = []
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            case_id = clean_inline_text(row[0] if row else "")
            if not case_id:
                continue

            test_cases.append(
                {
                    "id": case_id,
                    "escenario_referenciado": clean_inline_text(
                        row[header_index["historiadeusuariorequerimientoescenario"]]
                    ),
                    "escenario_referenciado_normalizado": normalize_lookup_text(
                        row[header_index["historiadeusuariorequerimientoescenario"]]
                    ),
                    "titulo": clean_inline_text(row[header_index["titulo"]]),
                    "descripcion": clean_inline_text(row[header_index["descripcion"]]),
                    "precondiciones": normalize_numbered_items(
                        row[header_index["precondiciones"]]
                    ),
                    "datos": normalize_bullet_items(row[header_index["datos"]]),
                    "pasos": normalize_numbered_items(row[header_index["pasos"]]),
                    "resultado_esperado": normalize_multiline_text(
                        row[header_index["resultadoesperado"]]
                    ),
                    "implementado": "No",
                }
            )

        if not test_cases:
            raise ValueError("No se encontraron casos validos en la hoja 'Casos de prueba'.")
        return test_cases
    finally:
        workbook.close()


def resolve_scenario_id(case_id: str, scenarios: OrderedDict[str, dict]) -> str | None:
    scenarios_by_slug = {scenario_id.casefold(): scenario_id for scenario_id in scenarios}
    candidate = case_id.strip()
    if candidate.casefold() in scenarios_by_slug:
        return scenarios_by_slug[candidate.casefold()]

    while "." in candidate:
        candidate = candidate.rsplit(".", 1)[0]
        if candidate.casefold() in scenarios_by_slug:
            return scenarios_by_slug[candidate.casefold()]
    return None


def longest_common_prefix_length(left: str, right: str) -> int:
    count = 0
    for left_char, right_char in zip(left.casefold(), right.casefold()):
        if left_char != right_char:
            break
        count += 1
    return count


def resolve_scenario_id_by_prefix(case_id: str, scenarios: OrderedDict[str, dict]) -> dict | None:
    prefix = case_id.strip()[:3]
    if len(prefix) < 3:
        return None

    candidates = [
        scenario_id for scenario_id in scenarios if scenario_id.casefold().startswith(prefix.casefold())
    ]
    if not candidates:
        return None

    selected = max(
        candidates,
        key=lambda scenario_id: (
            longest_common_prefix_length(case_id, scenario_id),
            -len(scenario_id),
            scenario_id.casefold(),
        ),
    )
    return {
        "prefix": prefix,
        "selected_scenario_id": selected,
        "candidate_scenario_ids": candidates,
    }


def normalize_lookup_text(value: object) -> str:
    text = normalize_scalar(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.casefold()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def assign_cases_to_scenarios(
    scenarios: OrderedDict[str, dict], test_cases: list[dict]
) -> tuple[list[dict], list[dict]]:
    unassigned: list[dict] = []
    prefix_fallback_matches: list[dict] = []
    scenarios_by_description = {
        scenario["descripcion_normalizada"]: scenario_id
        for scenario_id, scenario in scenarios.items()
        if scenario["descripcion_normalizada"]
    }

    for case in test_cases:
        scenario_id = resolve_scenario_id(case["id"], scenarios)
        if scenario_id is None:
            prefix_match = resolve_scenario_id_by_prefix(case["id"], scenarios)
            if prefix_match is not None:
                scenario_id = prefix_match["selected_scenario_id"]
                prefix_fallback_matches.append(
                    {
                        "caso_id": case["id"],
                        "prefijo": prefix_match["prefix"],
                        "escenario_asignado_id": scenario_id,
                        "escenario_asignado_descripcion": scenarios[scenario_id]["descripcion"],
                        "escenario_referenciado": case["escenario_referenciado"],
                        "candidatos": prefix_match["candidate_scenario_ids"],
                    }
                )
        if scenario_id is None:
            scenario_id = scenarios_by_description.get(case["escenario_referenciado_normalizado"])
        if scenario_id is None:
            unassigned.append(case)
            continue
        scenarios[scenario_id]["casos"].append(case)
    return unassigned, prefix_fallback_matches


def generate_prefix_fallback_report(prefix_fallback_matches: list[dict]) -> str:
    lines = [
        "# Reporte de asociaciones por prefijo",
        "",
        "Este reporte registra los casos de prueba cuyo escenario no se pudo resolver por ID exacto",
        "o por recorte de sufijos, y que fueron asociados usando el prefijo de 3 caracteres del ID del caso.",
        "",
        f"- Casos asociados por prefijo: {len(prefix_fallback_matches)}",
    ]

    if not prefix_fallback_matches:
        lines.extend(["", "_No se requirieron asociaciones por prefijo._"])
        return "\n".join(lines).rstrip() + "\n"

    for match in prefix_fallback_matches:
        lines.extend(
            [
                "",
                f"## Caso de prueba: {match['caso_id']}",
                f"- Prefijo usado: `{match['prefijo']}`",
                f"- Escenario asignado: `{match['escenario_asignado_id']}`",
                f"- Descripcion del escenario asignado: {match['escenario_asignado_descripcion']}",
                f"- Escenario reportado en el caso: {match['escenario_referenciado'] or '_vacio_'}",
                "- Escenarios candidatos con el mismo prefijo: "
                + ", ".join(f"`{candidate}`" for candidate in match["candidatos"]),
            ]
        )

    return "\n".join(lines).rstrip() + "\n"


def render_numbered_section(title: str, items: list[str]) -> list[str]:
    lines = [f"**{title}:**"]
    if items:
        lines.extend(f"{index}. {item}" for index, item in enumerate(items, start=1))
    else:
        lines.append("1. ")
    return lines


def render_bullet_section(title: str, items: list[str]) -> list[str]:
    lines = [f"**{title}:**"]
    if items:
        lines.extend(f"- {item}" for item in items)
    else:
        lines.append("- ")
    return lines


def render_test_case(case: dict) -> str:
    lines = [
        f"### Caso de prueba: {case['id']}",
        f"- Implementado: {case['implementado']}",
        "",
        f"**Titulo:** {case['titulo']}",
        "",
        f"**Descripcion:** {case['descripcion']}",
        "",
    ]
    lines.extend(render_numbered_section("Precondiciones", case["precondiciones"]))
    lines.append("")
    lines.extend(render_bullet_section("Datos", case["datos"]))
    lines.append("")
    lines.extend(render_numbered_section("Pasos", case["pasos"]))
    lines.append("")
    lines.append("**Resultado esperado:**")
    lines.append(case["resultado_esperado"])
    return "\n".join(lines).rstrip()


def generate_markdown(
    scenarios: OrderedDict[str, dict], unassigned_cases: list[dict]
) -> str:
    sections: list[str] = ["# Catalogo de casos de prueba"]

    for scenario in scenarios.values():
        scenario_lines = [
            f"## Escenario: {scenario['id']}",
            f"**Descripcion:** {scenario['descripcion']}",
            "",
        ]

        if scenario["casos"]:
            scenario_lines.extend(render_test_case(case) for case in scenario["casos"])
        else:
            scenario_lines.append("_Sin casos de prueba asociados._")

        sections.append("\n".join(scenario_lines).rstrip())

    if unassigned_cases:
        unassigned_lines = ["## Casos sin escenario asociado", ""]
        unassigned_lines.extend(render_test_case(case) for case in unassigned_cases)
        sections.append("\n".join(unassigned_lines).rstrip())

    return "\n\n---\n\n".join(sections) + "\n"


def validate_markdown(markdown: str) -> None:
    required_markers = [
        "# Catalogo de casos de prueba",
        "## Escenario:",
        "### Caso de prueba:",
        "- Implementado: No",
    ]
    missing = [marker for marker in required_markers if marker not in markdown]
    if missing:
        raise ValueError(
            "El Markdown generado no cumple con la estructura minima. "
            f"Faltan: {', '.join(missing)}"
        )


def build_catalog_payload(
    workbook_path: Path,
) -> tuple[OrderedDict[str, dict], list[dict], list[dict], list[dict]]:
    scenarios = load_scenarios(workbook_path)
    test_cases = load_test_cases(workbook_path)
    unassigned_cases, prefix_fallback_matches = assign_cases_to_scenarios(scenarios, test_cases)
    return scenarios, test_cases, unassigned_cases, prefix_fallback_matches


def build_cases_index(scenarios: OrderedDict[str, dict]) -> OrderedDict[str, dict]:
    cases_index: OrderedDict[str, list[dict]] = OrderedDict()

    for scenario in scenarios.values():
        for case in scenario["casos"]:
            entry = {
                "scenario_id": scenario["id"],
                "title": case["titulo"],
                "automation_readiness_hint": None,
                "source_case_ref": case["id"],
            }
            cases_index.setdefault(case["id"], []).append(entry)

    return cases_index


def build_json_catalog(
    workbook_path: Path,
    scenarios: OrderedDict[str, dict],
    unassigned_cases: list[dict],
    total_cases: int,
) -> OrderedDict[str, object]:
    json_scenarios: list[dict] = []

    for scenario in scenarios.values():
        json_scenarios.append(
            {
                "scenario_id": scenario["id"],
                "description": scenario["descripcion"],
                "cases": [
                    {
                        "case_id": case["id"],
                        "catalog_implemented_flag": case["implementado"],
                        "title": case["titulo"],
                        "description": case["descripcion"],
                        "preconditions": case["precondiciones"],
                        "data": case["datos"],
                        "steps": case["pasos"],
                        "expected_result": case["resultado_esperado"],
                        "source_case_ref": case["id"],
                    }
                    for case in scenario["casos"]
                ],
            }
        )

    payload: OrderedDict[str, object] = OrderedDict(
        [
            ("source_file", workbook_path.name),
            ("generated_at", datetime.now(timezone.utc).isoformat()),
            ("total_cases", total_cases),
            ("visible_scenarios_count", len(json_scenarios)),
            ("assigned_to_visible_scenarios_count", sum(len(scenario["cases"]) for scenario in json_scenarios)),
            ("unassigned_to_visible_scenarios_count", len(unassigned_cases)),
            ("scenarios", json_scenarios),
            ("cases_index", build_cases_index(scenarios)),
        ]
    )

    if unassigned_cases:
        payload["unassigned_cases"] = [
            {
                "case_id": case["id"],
                "catalog_implemented_flag": case["implementado"],
                "title": case["titulo"],
                "description": case["descripcion"],
                "preconditions": case["precondiciones"],
                "data": case["datos"],
                "steps": case["pasos"],
                "expected_result": case["resultado_esperado"],
                "source_case_ref": case["id"],
                "scenario_id": None,
            }
            for case in unassigned_cases
        ]

    return payload


def build_catalog(workbook_path: Path) -> tuple[dict, str, str, int, int, int]:
    scenarios, test_cases, unassigned_cases, prefix_fallback_matches = build_catalog_payload(
        workbook_path
    )
    json_catalog = build_json_catalog(workbook_path, scenarios, unassigned_cases, len(test_cases))
    markdown = generate_markdown(scenarios, unassigned_cases)
    prefix_report = generate_prefix_fallback_report(prefix_fallback_matches)
    validate_markdown(markdown)
    return (
        json_catalog,
        markdown,
        prefix_report,
        len(test_cases),
        len(unassigned_cases),
        len(prefix_fallback_matches),
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convierte un Excel de casos de prueba en un catalogo JSON y, opcionalmente, Markdown."
    )
    parser.add_argument("ruta_excel", type=Path, help="Ruta al archivo Excel de entrada.")
    parser.add_argument(
        "--ruta-salida-json",
        type=Path,
        default=Path(DEFAULT_JSON_OUTPUT),
        help=f"Ruta del archivo JSON de salida. Por defecto: {DEFAULT_JSON_OUTPUT}",
    )
    parser.add_argument(
        "--ruta-salida-markdown",
        type=Path,
        default=None,
        help=(
            "Ruta del archivo Markdown de salida. "
            f"Si se omite, no se genera Markdown. Sugerido: {DEFAULT_MARKDOWN_OUTPUT}"
        ),
    )
    parser.add_argument(
        "--ruta-reporte-prefijo",
        type=Path,
        default=Path(DEFAULT_PREFIX_REPORT),
        help=(
            "Ruta del reporte Markdown para asociaciones resueltas por prefijo. "
            f"Por defecto: {DEFAULT_PREFIX_REPORT}"
        ),
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    workbook_path = args.ruta_excel
    json_output_path = args.ruta_salida_json
    markdown_output_path = args.ruta_salida_markdown
    prefix_report_path = args.ruta_reporte_prefijo

    if not workbook_path.exists():
        print(f"El archivo Excel no existe: {workbook_path}", file=sys.stderr)
        return 1

    try:
        (
            json_catalog,
            markdown,
            prefix_report,
            total_cases,
            unassigned_cases,
            prefix_assigned_cases,
        ) = build_catalog(workbook_path)
        json_output_path.write_text(
            json.dumps(json_catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        if markdown_output_path is not None:
            markdown_output_path.write_text(markdown, encoding="utf-8")
        prefix_report_path.write_text(prefix_report, encoding="utf-8")
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        return 1

    markdown_message = (
        f" Markdown generado en '{markdown_output_path}'."
        if markdown_output_path is not None
        else ""
    )
    print(
        f"Catalogo JSON generado en '{json_output_path}' con {total_cases} casos "
        f"y {unassigned_cases} casos sin escenario asociado."
        f"{markdown_message} "
        f"Reporte de prefijo generado en '{prefix_report_path}' con "
        f"{prefix_assigned_cases} asociaciones."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
